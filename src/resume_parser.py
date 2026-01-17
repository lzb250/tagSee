# src/resume_parser.py
import os
import re
from pathlib import Path
from typing import Union, Optional
from .cross_platform_utils import safe_path, normalize_line_endings

try:
    import pdfplumber
    PDF_AVAILABLE = True
except ImportError:
    PDF_AVAILABLE = False

try:
    import docx
    DOCX_AVAILABLE = True
except ImportError:
    DOCX_AVAILABLE = False

class ResumeParser:
    """跨平台简历文本提取器"""

    def __init__(self):
        self.supported_extensions = {'.txt', '.pdf', '.docx'}

    def extract_text(self, file_path: Union[str, Path]) -> Optional[str]:
        """
        从不同格式的文件中提取文本
        支持: .txt, .pdf, .docx
        """
        file_path = safe_path(file_path)

        if not file_path.exists():
            raise FileNotFoundError(f"文件不存在: {file_path}")

        file_extension = file_path.suffix.lower()

        if file_extension not in self.supported_extensions:
            raise ValueError(f"不支持的文件格式: {file_extension}. 支持: {self.supported_extensions}")

        try:
            if file_extension == '.txt':
                return self._extract_from_txt(file_path)
            elif file_extension == '.pdf':
                return self._extract_from_pdf(file_path)
            elif file_extension == '.docx':
                return self._extract_from_docx(file_path)

        except Exception as e:
            raise RuntimeError(f"处理文件 {file_path} 时出错: {str(e)}")

    def _extract_from_txt(self, file_path: Path) -> str:
        """从文本文件提取内容"""
        encodings = ['utf-8', 'gbk', 'gb2312', 'latin1']

        for encoding in encodings:
            try:
                with open(file_path, 'r', encoding=encoding) as f:
                    content = f.read()
                return normalize_line_endings(content)
            except UnicodeDecodeError:
                continue

        # 如果所有编码都失败，使用错误处理
        with open(file_path, 'r', encoding='utf-8', errors='replace') as f:
            return normalize_line_endings(f.read())

    def _extract_from_pdf(self, file_path: Path) -> str:
        """从PDF文件提取内容"""
        if not PDF_AVAILABLE:
            raise ImportError("需要安装 pdfplumber 来处理 PDF 文件: pip install pdfplumber")

        text = ""
        with pdfplumber.open(file_path) as pdf:
            for page in pdf.pages:
                page_text = page.extract_text()
                if page_text:
                    text += page_text + "\n"

        return normalize_line_endings(text)

    def _extract_from_docx(self, file_path: Path) -> str:
        """从DOCX文件提取内容"""
        if not DOCX_AVAILABLE:
            raise ImportError("需要安装 python-docx 来处理 DOCX 文件: pip install python-docx")

        doc = docx.Document(file_path)
        paragraphs = []
        for paragraph in doc.paragraphs:
            if paragraph.text.strip():
                paragraphs.append(paragraph.text)

        # 提取表格内容（如果有的话）
        for table in doc.tables:
            for row in table.rows:
                for cell in row.cells:
                    if cell.text.strip():
                        paragraphs.append(cell.text)

        return normalize_line_endings("\n".join(paragraphs))

    def clean_text(self, text: str) -> str:
        """清理文本"""
        if not text:
            return ""

        # 移除多余的空白字符
        text = re.sub(r'\s+', ' ', text)

        # 移除特殊控制字符，但保留中文、英文、数字和常用符号
        text = re.sub(r'[^\u4e00-\u9fff\w\s\-+.,;:!?@#$%^&*()_+=\[\]{}|\\:"<>?/~`]', ' ', text)

        return text.strip()

# 使用示例
if __name__ == "__main__":
    parser = ResumeParser()
    # 测试代码...